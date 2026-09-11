"""Tests for src/reports/excel.py."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.database.connection import connect_and_init
from src.database.repositories import (
    imports_repo,
    issues_repo,
    order_items_repo,
    orders_repo,
    payments_repo,
)
from src.reports import excel


@pytest.fixture()
def conn(tmp_path: Path):
    connection = connect_and_init(tmp_path / "test.db")
    yield connection
    connection.close()


@pytest.fixture()
def import_id(conn) -> int:
    return imports_repo.create_import(conn, "x.csv", "mixed", "h-mix")


def _payment(order_id: str, **overrides) -> dict:
    base = {
        "payment_id": f"p-{order_id}",
        "order_id": order_id,
        "gross_amount": 100.0,
        "fees": 10.0,
        "net_amount": 90.0,
        "refund_amount": 0.0,
        "currency": "USD",
        "order_date": "2026-01-15",
        "funds_available": "2026-01-18",
        "status": "SETTLED",
    }
    base.update(overrides)
    return base


def _item(order_id: str, txn: str, listing: str, **overrides) -> dict:
    base = {
        "transaction_id": txn,
        "order_id": order_id,
        "sale_date": "2026-01-15",
        "listing_id": listing,
        "item_name": f"Item {listing}",
        "quantity": 1,
        "price": 100.0,
        "item_total": 100.0,
        "currency": "USD",
        "sku": None,
    }
    base.update(overrides)
    return base


def _populate(conn, import_id) -> None:
    payments_repo.insert_payments(
        conn,
        [
            _payment("1001", gross_amount=120.0, fees=12.0, net_amount=108.0),
            _payment("1002", gross_amount=80.0, fees=8.0, net_amount=72.0),
        ],
        import_id,
    )
    order_items_repo.insert_order_items(
        conn,
        [
            _item("1001", "tx1", "L1", item_total=120.0),
            _item("1002", "tx2", "L2", item_total=80.0),
        ],
        import_id,
    )


def test_report_file_is_created(conn, import_id, tmp_path: Path) -> None:
    _populate(conn, import_id)
    out = tmp_path / "report.xlsx"
    result = excel.build_report(conn, out)
    assert result == out
    assert out.exists()
    assert out.stat().st_size > 0


def test_report_has_expected_sheets(conn, import_id, tmp_path: Path) -> None:
    _populate(conn, import_id)
    out = tmp_path / "report.xlsx"
    excel.build_report(conn, out)

    wb = load_workbook(out, read_only=True)
    try:
        expected = {
            "Summary", "Sales", "Products",
            "Countries", "Data Quality", "Insights",
        }
        assert expected.issubset(set(wb.sheetnames))
    finally:
        wb.close()


def test_summary_sheet_contains_kpi_labels(conn, import_id, tmp_path: Path) -> None:
    _populate(conn, import_id)
    out = tmp_path / "report.xlsx"
    excel.build_report(conn, out)

    wb = load_workbook(out, read_only=True)
    try:
        ws = wb["Summary"]
        values = [cell.value for row in ws.iter_rows() for cell in row]
        assert "Gross Revenue" in values
        assert "Net Revenue" in values
        assert "Average Order Value" in values
    finally:
        wb.close()


def test_products_sheet_contains_listing_ids(conn, import_id, tmp_path: Path) -> None:
    _populate(conn, import_id)
    out = tmp_path / "report.xlsx"
    excel.build_report(conn, out)

    wb = load_workbook(out, read_only=True)
    try:
        ws = wb["Products"]
        values = [cell.value for row in ws.iter_rows() for cell in row]
        assert "L1" in values
        assert "L2" in values
    finally:
        wb.close()


def test_data_quality_sheet_reflects_issues(conn, import_id, tmp_path: Path) -> None:
    _populate(conn, import_id)
    issues_repo.add_issue(
        conn, "test_check", "WARNING", "sample warning", table_name="orders"
    )
    out = tmp_path / "report.xlsx"
    excel.build_report(conn, out)

    wb = load_workbook(out, read_only=True)
    try:
        ws = wb["Data Quality"]
        values = [cell.value for row in ws.iter_rows() for cell in row]
        assert "sample warning" in values
    finally:
        wb.close()


def test_run_reconciliation_flag_populates_issues(
    conn, import_id, tmp_path: Path
) -> None:
    # Payment without a matching order → ERROR issue.
    payments_repo.insert_payments(conn, [_payment("9999")], import_id)
    out = tmp_path / "report.xlsx"
    excel.build_report(conn, out, run_reconciliation=True)

    wb = load_workbook(out, read_only=True)
    try:
        ws = wb["Data Quality"]
        values = [cell.value for row in ws.iter_rows() for cell in row]
        assert "payment_without_order" in values
    finally:
        wb.close()


def test_creates_parent_directory(conn, import_id, tmp_path: Path) -> None:
    _populate(conn, import_id)
    nested = tmp_path / "a" / "b" / "report.xlsx"
    excel.build_report(conn, nested)
    assert nested.exists()